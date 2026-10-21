# Se ejecuta dentro del espacio global de app.py para conservar compatibilidad.
import streamlit as st
import pandas as pd
import os
from io import BytesIO

st.subheader("🚚 CREAR GUIA RUTA LA PAZ")

# --------------------- SIDEBAR MENU ---------------------
st.sidebar.title("Menú de Opciones")
menu = st.sidebar.radio("Ir a:", [
    "🌎 Detalle General",
    "🏷️ Filtrar Solicitudes",
    "🧩 Combinar Archivos",
    "🔎 Generar Sintaxis de Búsqueda",
    "📦 Preparar Ruta Almacén La Paz",
    "📦 Preparar Guias de Ruta",
    "🗑️ Eliminar Archivos"

])

#----------------------- 0. DETALLE GENERAL -----------------------
def detalle_general():
    st.title("🧮 SOLICITUDES POR REGIONALES")

    # ==============================
    # CARGA ARCHIVO
    # ==============================
    co1,co2 = st.columns([15.04, 8.05])
    with co1:
        uploaded_file = st.file_uploader(
            "📁 CARGAR ARCHIVO EXCEL",
            type=["csv", "xlsx"]
        )

    # ==============================
    # TIPOS DE SOLICITUD
    # ==============================
    TIPOS_SOLICITUD = [
        "SE1 - Envío de Materiales",
        "SR1 - Recojo de items nuevos (Cajas o Files)",
        "SR2 - Recojo de Items por devolución (Cajas o Files)"
    ]

    # ==============================
    # CENTROS CENTRALIZADOS
    # ==============================
    CENTROS_CENTRALIZADO = {
        "REGIONAL LA PAZ": {
            "ZONA SUR": ["212 - COTA COTA", "208 - SAN MIGUEL", "219 - OBRAJES"],
            "ZONA ESTE": [
                "204 - MIRAFLORES", 
                "218 - VILLA ARMONIA", 
                "211 - CRUCE VILLA COPACABANA", 
                "223 - PAMPAHASI"],
            "CENTRO": [
                "221 - LA PORTADA", 
                "205 - EL TEJAR", 
                "216 - GRAN PODER", 
                "202 - GARITA", "206 - ALONSO DE MENDOZA",
                "213 - NORMALIZADORA LA PAZ", 
                "214 - SOL AMIGO LA PAZ", 
                "220 - TEMBLADERANI", 
                "201 - SAN PEDRO",
                "295 - OFICINA NACIONAL", 
                "210 - CAMACHO", 
                "209 - BALLIVIAN", 
                "200 - REGIONAL LA PAZ"],
            "NORTE": [
                "203 - VILLA FATIMA", 
                "224 - CHUQUIAGUILLO", 
                "222 - PERIFERICA", 
                "217 - VINO TINTO"]
        },
        "REGIONAL EL ALTO": {
            "NORTE": [
                "253 - RIO SECO", 
                "276 - VILLA INGENIO", 
                "272 - SAN ROQUE", 
                "266 - MERCADO EL CARMEN RÍO SECO",
                "275 - FRANZ TAMAYO", 
                "251 - 16 DE JULIO", 
                "279 - CHACALTAYA", 
                "277 - FERROPETROL",
                "280 - LAGUNAS EX PARADA 8"],
            "CENTRO": [
                "270 - 12 DE OCTUBRE", 
                "252 - LA CEJA", 
                "263 - VILLA DOLORES", 
                "258 - NORMALIZACIÓN EL ALTO",
                "262 - SATELITE",
                "264 - SOL AMIGO EL ALTO",
                "265 - AGENCIA MOVIL",  
                "250 - REGIONAL EL ALTO"],
            "OESTE": [
                "254 - VILLA ADELA", 
                "261 - BOLIVIA", 
                "274 - QUISWARAS"],
            "SUR": [
                "267 - SANTIAGO II", 
                "273 - EL KENKO", 
                "260 - SENKATA", 
                "269 - VENTILLA"],
            "VIACHA": [
                "271 - INGAVI", 
                "256 - VIACHA"]
        },
        "REGIONAL ORURO": {
            "ORURO": [
						"400 - REGIONAL ORURO",
                "401 - CENTRAL", 
                "407 - NORMALIZACION ORURO", 
                "408 - PUNTO AMIGO ORURO", 
                "409 - VIRGEN DEL SOCAVÓN",
                "410 - TAGARETE", 
                "411 - TACNA"]
        },
        "REGIONAL SUCRE": {
            "SUCRE": [
                "100 - REGIONAL CHUQUISACA", 
                "101 - MERCADO CAMPESINO SUCRE", 
                "102 - NORMALIZACION SUCRE", 
                "103 - ESPAÑA",
                "104 - SOL AMIGO SUCRE", 
                "105 - GERMAN MENDOZA", 
                "106 - CHARCAS", 
                "107 - ZUDAÑEZ",
                "108 - LAS AMERICAS", 
                "109 - LAJASTAMBO"]
        },
        "REGIONAL TARIJA": {
            "": [
						"600 - REGIONAL TARIJA",
                "601 - MERCADO CAMPESINO TARIJA",
                "603 - SUR",
                "604 - SOL AMIGO TARIJA",
                "605 - 15 DE ABRIL",
                "606 - NORMALIZACIÓN TARIJA",
                "607 - TABLADITA",
                "608 - YACUIBA",
                "609 - PALMARCITO",
                "610 - MERCADO CAMPESINO YACUIBA"]
        },

        "BANCO NACIONAL DE BOLIVIA": {
            "BNB - LA PAZ": [
                "130 - BANCO NACIONAL DE BOLIVIA - LA PAZ"],
            "BNB - COCHABAMBA":[
                "132 - BANCO NACIONAL DE BOLIVIA - COCHABAMBA"],
            "BNB - TARIJA":[
                "136 - BANCO NACIONAL DE BOLIVIA - TARIJA"],
            "BNB - ORURO":[
                "134 - BANCO NACIONAL DE BOLIVIA - ORURO"],
            "BNB - SUCRE":[
                "133 - BANCO NACIONAL DE BOLIVIA - SUCRE"],
            "BNB - SANTA CRUZ":[
                "131 - BANCO NACIONAL DE BOLIVIA - SANTA CRUZ"],
            "BNB - BENI":[
                "137 - BANCO NACIONAL DE BOLIVIA - BENI"],
            "BNB - POTOSI":[
                "135 - BANCO NACIONAL DE BOLIVIA - POTOSI"],
            "BNB - PANDO":[
                "138 - BANCO NACIONAL DE BOLIVIA - PANDO"],
            "BNB - POLYSISTEMAS":[
                "C008 - POLYSISTEMAS"]

        },

        "BANCO FIE": {
            "BFIE LA PAZ": [
                "124-1 - BFIE LA PAZ",
						"REGIONAL LA PAZ - CORIPATA",
						"REGIONAL LA PAZ - EL TEJAR",
						"REGIONAL LA PAZ - PALOS BLANCOS",
						"REGIONAL LA PAZ - LA PORTADA",
						"REGIONAL LA PAZ - VILLA FATIMA",
						"REGIONAL LA PAZ - PLAZA EGUINO",
						"REGIONAL LA PAZ - VILLA COPACABANA",
						"REGIONAL LA PAZ - SOPOCACHI",
						"REGIONAL LA PAZ - CHASQUIPAMPA",
						"REGIONAL LA PAZ - GRAN PODER",
						"REGIONAL LA PAZ - BUENOS AIRES",
						"REGIONAL LA PAZ - MIRAFLORES",
						"REGIONAL LA PAZ - TEMBLADERANI",
						"REGIONAL LA PAZ - SAN PEDRO",
						"PAMPAHASI",
						"REGIONAL LA PAZ - PAMPAHASI",
						"BUENOS AIRES",
						"REGIONAL LA PAZ - BUENOS AIRES",
						"REGIONAL LA PAZ - COROICO",
						"REGIONAL LA PAZ - RURRENABAQUE",
						"REGIONAL LA PAZ - PERIFERICA",
						"REGIONAL LA PAZ - ZONA SUR",
						"REGIONAL LA PAZ - CAMACHO",
						"REGIONAL LA PAZ - SAN MIGUEL",
						"REGIONAL LA PAZ - COBIJA",
						"REGIONAL LA PAZ - CHULUMANI",
						"REGIONAL LA PAZ - CARANAVI"
						
					],
					
					"BFIE EL ALTO":[
						"REGIONAL EL ALTO - SENKATA",
						"REGIONAL EL ALTO - 21 DE OCTUBRE",
						"REGIONAL EL ALTO - ACHOCALLA",
						"REGIONAL EL ALTO - NUEVOS HORIZONTES",
						"REGIONAL EL ALTO - LIBERTAD",
						"REGIONAL EL ALTO - PATACAMAYA",
						"REGIONAL EL ALTO - TERMINAL",
						"REGIONAL EL ALTO - VILLA YUNGUYO",
						"REGIONAL EL ALTO - BALLIVIAN",
						"REGIONAL EL ALTO - DESAGUADERO",
						"REGIONAL EL ALTO - ACHACACHI",
						"REGIONAL EL ALTO - COPACABANA",
						"REGIONAL EL ALTO - VILLA ADELA",
						"REGIONAL EL ALTO - VILLA ESPERANZA",
						"REGIONAL EL ALTO - OFICINA CENTRAL",
						"REGIONAL EL ALTO - NUEVO AMANECER",
						"REGIONAL EL ALTO - PANAMERICANA",
						"REGIONAL EL ALTO - RIO SECO",
						"REGIONAL EL ALTO - LA CEJA",
						"REGIONAL EL ALTO - ROMERO PAMPA",
						"REGIONAL EL ALTO - VILLA DOLORES",
						"REGIONAL EL ALTO - VENTILLA",
						"REGIONAL EL ALTO - 16 DE JULIO",
						"REGIONAL EL ALTO - VENTILLA",
						"REGIONAL EL ALTO - VIACHA",
						"REGIONAL EL ALTO - 12 DE OCTUBRE",
						"REGIONAL EL ALTO - SATELITE",
						"REGIONAL EL ALTO - SANTIAGO",
						"REGIONAL EL ALTO - PACAJES"
						
					],

					"BFIE COCHABAMBA": [
						"REGIONAL COCHABAMBA - VINTO",
						"REGIONAL COCHABAMBA - TAMBORADA",
						"REGIONAL COCHABAMBA - CRUCE TAQUIÑA",
						"REGIONAL COCHABAMBA - QUILLACOLLO",
						"REGIONAL COCHABAMBA - SACABA",
						"REGIONAL COCHABAMBA - JORDAN",
						"REGIONAL COCHABAMBA - COLCAPIRHUA",
						"REGIONAL COCHABAMBA - PUNATA",
						"REGIONAL COCHABAMBA - PACATA",
						"REGIONAL COCHABAMBA - CALA CALA",
						"REGIONAL COCHABAMBA - KANATA",
						"REGIONAL COCHABAMBA - SAN MARTIN",
						"REGIONAL COCHABAMBA - LA CANCHA",
						"REGIONAL COCHABAMBA - NATANIEL AGUIRRE",
						"REGIONAL COCHABAMBA - HEROINAS",
						"REGIONAL COCHABAMBA - PETROLERA",
						"REGIONAL COCHABAMBA - ENTRE RIOS",
						"REGIONAL COCHABAMBA - ALALAY",
						"REGIONAL COCHABAMBA - IVIRGARZAMA",
						"REGIONAL COCHABAMBA - CHIMORE",
						"REGIONAL COCHABAMBA - URKUPIÑA",
						"REGIONAL COCHABAMBA - CRUCE TAQUIÑA"

					],

					"BFIE ORURO":[
						"REGIONAL ORURO - PAGADOR",
						"REGIONAL ORURO - SUCURSAL ORURO",
						"REGIONAL ORURO - ESPAÑA",
						"REGIONAL ORURO - LLALLAGUA",
						"REGIONAL ORURO - BOLIVAR"
					],

					"BFIE TARIJA":[
						"REGIONAL TARIJA - YACUIBA",
						"REGIONAL TARIJA - CAMARGO",
						"REGIONAL TARIJA - LA TABLADA",
						"REGIONAL TARIJA - ANDALUCIA",
						"REGIONAL TARIJA - LUIS DE FUENTE",
						"REGIONAL TARIJA - GUADALUPE",
						"REGIONAL TARIJA - BERMEJO",
						"REGIONAL TARIJA - GUADALQUIVIR",
						"REGIONAL TARIJA - VALLE DE CONCEPCION",
						"REGIONAL TARIJA - ARANJUEZ",
						"REGIONAL TARIJA - VILLAMONTES"

					],

					"BFIE SUCRE":[
						"REGIONAL SUCRE - PADILLA",
						"REGIONAL SUCRE - 25 DE MAYO",
						"REGIONAL SUCRE - CHARCAS",
						"REGIONAL SUCRE - MONTEAGUDO",
						"REGIONAL SUCRE - MESA VERDE"
					],

					"BFIE POTOSI":[
						"REGIONAL POTOSI - SAN ROQUE",
						"REGIONAL POTOSI - CERRO DE PLATA",
						"REGIONAL POTOSI - COTAGAITA",
						"REGIONAL POTOSI - TUPIZA",
						"REGIONAL POTOSI - 10 DE NOVIEMBRE",
						"REGIONAL POTOSI - UYUNI",
						"REGIONAL POTOSI - BETANZOS",
						"REGIONAL POTOSI - NUEVA TERMINAL",
						"REGIONAL POTOSI - JUNIN",
						"REGIONAL POTOSI - VILLAZON"

					],
					
            "BFIE SANTA CRUZ": [
                "124-2 - BFIE SANTA CRUZ"],
            "BFIE INTERNO": [   
                "C008 - POLYSISTEMAS"]
                                    
        },

        "JTI BOLIVIA": {
            "JTI EXTERNO": [
                "367 - JTI BOLIVIA"],
            "JTI INTERNO": [   
                "C008 - POLYSISTEMAS"]
        },

        "REGIONAL COCHABAMBA": {
            "": [
                "301 - ESTEBAN ARCE",
                "302 - SAN MARTIN",
                "303 - HUAYRA KHASA",
                "305 - CRUCE TAQUIÑA",
                "306 - QUILLACOLLO",
                "307 - COLCA PIRHUA",
                "309 - MUYURINA",
                "310 - NORMALIZACIÓN COCHABAMBA",
                "311 - RECAUDADORA JORDAN",
                "312 - SACABA",
                "313 - VILLA GALINDO",
                "314 - PUNATA",
                "316 - AYACUCHO",
                "317 - SOL AMIGO COCHABAMBA",
                "318 - PANAMERICANA",
                "320 - CLIZA",
                "321 - VINTO",
                "322 - REPÚBLICA",
                "323 - TIQUIPAYA",
                "324 - QUINTANILLA",
                "325 - JORDAN",
                "326 - PLAZA BOLIVAR",
                "327 - PETROLERA",
                "328 - LA CHIMBA",
                "329 - AMERICA",
                "331 - EL AVION",
                "332 - VILLA PAGADOR",
                "333 - PACATA"]
        },
        "REGIONAL BENI": {
            "": [
                "801 - TRINIDAD",
                "802 - RIBERALTA",
                "803 - GUAYARAMERIN"]
        },
        "REGIONAL COBIJA": {
            "": [
                "901 - COBIJA",
                "902 - TAJIBOS",
                "903 - NORMALIZACIÓN"]
        },
        "REGIONAL POTOSI": {
            "": [
						"500 - REGIONAL POTOSI", 
                "501 - MERCADO UYUNI",
                "502 - SOL AMIGO POTOSI",
                "503 - BOULEVARD",
					    "505 - BANDERAS", 
					    "506 - MURILLO"] 
        },
        "REGIONAL SANTA CRUZ": {
            "": [
                "701 - CASCO VIEJO",
                "702 - EL PARI",
                "703 - MUTUALISTA",
                "704 - 1RO. DE MAYO",
                "705 - MONTERO",
                "706 - EL TORNO",
                "709 - PIRAI",
                "711 - PLAN 3000",
                "715 - LA GUARDIA",
                "716 - ALTO SAN PEDRO",
                "718 - NORTE",
                "719 - SOL AMIGO SANTA CRUZ",
                "721 - ARROYO CONCEPCIÓN",
                "723 - PAMPA DE LA ISLA",
                "724 - COLORADA",
                "725 - Minero",
                "726 - SAN JULIAN",
                "727 - SAN JOSE",
                "728 - 2 DE AGOSTO",
                "730 - LOS LOTES",
                "731 - GERMAN MORENO",
                "732 - YAPACANI",
                "734 - SATELITE NORTE",
                "735 - NORTE I",
                "736 - EL QUIOR",
                "737 - VIRGEN DE LUJAN",
                "738 - EL BAJIO",
                "739 - CRISTO REDENTOR"]
        },
        "MOLINO ANDINO" :{
            "MOLINO ANDINO" : [
                "620 - MOLINO ANDINO S.A."]
        },
        "ROCHE BOLIVIA" :{
            "ROCHE BOLIVIA" :[
                "599 - ROCHE BOLIVIA S.R.L."]
        }
                        
    }

    # ==============================
    # PROCESAMIENTO
    # ==============================
    if uploaded_file:

        col1,col2 = st.columns([15.04, 8.05])
        with col1:
                
            if uploaded_file.name.endswith(".csv"):
                df = pd.read_csv(uploaded_file)
            else:
                df = pd.read_excel(uploaded_file)

            df.columns = df.columns.str.strip()

            if "Centro de Costo" in df.columns:
                df["Centro de Costo"] = (
                    df["Centro de Costo"]
                    .astype(str)
                    .str.extract(r'(\d+)')
                )

            if "Cantidad" in df.columns:
                df["Cantidad"] = pd.to_numeric(df["Cantidad"], errors="coerce").fillna(0)

            if "Tipo de Solicitud" in df.columns:
                df["Tipo de Solicitud"] = df["Tipo de Solicitud"].astype(str).str.strip()

            st.subheader("FILTRAR TIPO DE SOLICITUD")

            tipos_seleccionados = st.multiselect(
                "Selecciona uno o varios tipos:",
                options=TIPOS_SOLICITUD,
                default=TIPOS_SOLICITUD
            )

            df_filtrado = df[df["Tipo de Solicitud"].isin(tipos_seleccionados)]

        st.markdown("## 🌎 DETALLE GENERAL")

        # ==============================
        # CREAR DOS COLUMNAS
        # ==============================
        col1, col2 = st.columns(2)

        regionales = list(CENTROS_CENTRALIZADO.items())
        mitad = len(regionales) // 2 + len(regionales) % 2

        lista1 = regionales[:mitad]
        lista2 = regionales[mitad:]

        # ==============================
        # FUNCIÓN PARA MOSTRAR REGIONAL
        # ==============================
        def mostrar_regionales(lista):

            for regional, subzonas in lista:

                total_regional = 0

                with st.expander(f"📍 {regional}", expanded=False):

                    for subzona, lista_centros in subzonas.items():

                        total_subzona = 0
                        detalle_centros = []

                        for centro in lista_centros:

                            codigo = centro.split("-")[0].strip()

                            total = df_filtrado[
                                df_filtrado["Centro de Costo"] == codigo
                            ]["Cantidad"].sum()

                            total_subzona += total
                            total_regional += total

                            detalle_centros.append({
                                "Centro": centro,
                                "Cantidad": int(total)
                            })

                        detalle_df = pd.DataFrame(detalle_centros)\
                            .sort_values(by="Cantidad", ascending=False)

                        with st.expander(
                            f"🔹 {subzona} | Total: {int(total_subzona)}",
                            expanded=False
                        ):
                            st.dataframe(detalle_df, use_container_width=True)

                    st.markdown(f"### 🧮 TOTAL REGIONAL: {int(total_regional)}")

        with col1:
            mostrar_regionales(lista1)

        with col2:
            mostrar_regionales(lista2)
                                
# --------------------- 1. FILTRAR SOLICITUDES V2 ---------------------
def filtrar_solicitudes():

    st.title("📊 FILTRADOR DE SOLICITUDES PARA ELABORAR RUTAS")

    co1,co2 = st.columns([15.04, 8.05])
    with co1:
        # Subida del archivo
        uploaded_file = st.file_uploader("📁 CARGAR ARCHIVO EXCEL", type=["csv","xlsx"])
    #with co2:
        # Input para ruta personalizada (opcional)
        #carpeta_guardado = st.text_input(
            #"📁 Ingresa la ruta para para guardar el Archivo Filtrado :",
            #placeholder="Ej: C:/Users/TuUsuario/Desktop"
        #)

    # Definiciones de filtros
    TIPOS_SOLICITUD = [
        "SE1 - Envío de Materiales",
        "SR1 - Recojo de items nuevos (Cajas o Files)",
        "SR2 - Recojo de Items por devolución (Cajas o Files)"
    ]

    # Diccionario de Centros de Costo
    CENTROS_CENTRALIZADO = {
        "REGIONAL LA PAZ": {
            "ZONA SUR": [
                "212 - COTA COTA", 
                "208 - SAN MIGUEL", 
                "219 - OBRAJES"
                ],
            "ZONA ESTE": [
                "204 - MIRAFLORES", 
                "218 - VILLA ARMONIA", 
                "211 - CRUCE VILLA COPACABANA", 
                "223 - PAMPAHASI"],
            "CENTRO": [
                "221 - LA PORTADA", 
                "205 - EL TEJAR", 
                "216 - GRAN PODER", 
                "202 - GARITA", "206 - ALONSO DE MENDOZA",
                "213 - NORMALIZADORA LA PAZ", 
                "214 - SOL AMIGO LA PAZ", 
                "220 - TEMBLADERANI", 
                "201 - SAN PEDRO",
                "295 - OFICINA NACIONAL", 
                "210 - CAMACHO", 
                "209 - BALLIVIAN", 
                "200 - REGIONAL LA PAZ"],
            "NORTE": [
                "203 - VILLA FATIMA", 
                "224 - CHUQUIAGUILLO", 
                "222 - PERIFERICA", 
                "217 - VINO TINTO"]
        },
        "REGIONAL EL ALTO": {
            "NORTE": [
                "253 - RIO SECO", 
                "276 - VILLA INGENIO", 
                "272 - SAN ROQUE", 
                "266 - MERCADO EL CARMEN RÍO SECO",
                "275 - FRANZ TAMAYO", 
                "251 - 16 DE JULIO", 
                "279 - CHACALTAYA", 
                "277 - FERROPETROL",
                "280 - LAGUNAS EX PARADA 8"],
            "CENTRO": [
                "270 - 12 DE OCTUBRE", 
                "252 - LA CEJA", 
                "263 - VILLA DOLORES", 
                "258 - NORMALIZACIÓN EL ALTO",
                "262 - SATELITE",
                "264 - SOL AMIGO EL ALTO",
                "265 - AGENCIA MOVIL",  
                "250 - REGIONAL EL ALTO"],
            "OESTE": [
                "254 - VILLA ADELA", 
                "261 - BOLIVIA", 
                "274 - QUISWARAS"],
            "SUR": [
                "267 - SANTIAGO II", 
                "273 - EL KENKO", 
                "260 - SENKATA", 
                "269 - VENTILLA"],
            "VIACHA": [
                "271 - INGAVI", 
                "256 - VIACHA"]
        },
        "REGIONAL ORURO": {
            "ORURO": [
						"400 - REGIONAL ORURO",
                "401 - CENTRAL", 
                "407 - NORMALIZACION ORURO", 
                "408 - PUNTO AMIGO ORURO", 
                "409 - VIRGEN DEL SOCAVÓN",
                "410 - TAGARETE", 
                "411 - TACNA"]
        },
        "REGIONAL SUCRE": {
            "SUCRE": [
                "100 - REGIONAL CHUQUISACA", 
                "101 - MERCADO CAMPESINO SUCRE", 
                "102 - NORMALIZACION SUCRE", 
                "103 - ESPAÑA",
                "104 - SOL AMIGO SUCRE", 
                "105 - GERMAN MENDOZA", 
                "106 - CHARCAS", 
                "107 - ZUDAÑEZ",
                "108 - LAS AMERICAS", 
                "109 - LAJASTAMBO"]
        },
        "REGIONAL TARIJA": {
            "": [
						"600 - REGIONAL TARIJA",
                "601 - MERCADO CAMPESINO TARIJA",
                "603 - SUR",
                "604 - SOL AMIGO TARIJA",
                "605 - 15 DE ABRIL",
                "606 - NORMALIZACIÓN TARIJA",
                "607 - TABLADITA",
                "608 - YACUIBA",
                "609 - PALMARCITO",
                "610 - MERCADO CAMPESINO YACUIBA"
            ]
        },

        "BANCO NACIONAL DE BOLIVIA": {
            "BNB - LA PAZ": [
                "130 - BANCO NACIONAL DE BOLIVIA - LA PAZ"],
            "BNB - COCHABAMBA":[
                "132 - BANCO NACIONAL DE BOLIVIA - COCHABAMBA"],
            "BNB - TARIJA":[
                "136 - BANCO NACIONAL DE BOLIVIA - TARIJA"],
            "BNB - ORURO":[
                "134 - BANCO NACIONAL DE BOLIVIA - ORURO"],
            "BNB - SUCRE":[
                "133 - BANCO NACIONAL DE BOLIVIA - SUCRE"],
            "BNB - SANTA CRUZ":[
                "131 - BANCO NACIONAL DE BOLIVIA - SANTA CRUZ"],
            "BNB - BENI":[
                "137 - BANCO NACIONAL DE BOLIVIA - BENI"],
            "BNB - POTOSI":[
                "135 - BANCO NACIONAL DE BOLIVIA - POTOSI"],
            "BNB - PANDO":[
                "138 - BANCO NACIONAL DE BOLIVIA - PANDO"],
            "BNB - POLYSISTEMAS":[
                "C008 - POLYSISTEMAS"]

        },

        "BANCO FIE": {
            "BFIE LA PAZ": [
                "124-1 - BFIE LA PAZ",
						"REGIONAL LA PAZ - CORIPATA",
						"REGIONAL LA PAZ - EL TEJAR",
						"REGIONAL LA PAZ - PALOS BLANCOS",
						"REGIONAL LA PAZ - LA PORTADA",
						"REGIONAL LA PAZ - VILLA FATIMA",
						"REGIONAL LA PAZ - PLAZA EGUINO",
						"REGIONAL LA PAZ - VILLA COPACABANA",
						"REGIONAL LA PAZ - SOPOCACHI",
						"REGIONAL LA PAZ - CHASQUIPAMPA",
						"REGIONAL LA PAZ - GRAN PODER",
						"REGIONAL LA PAZ - BUENOS AIRES",
						"REGIONAL LA PAZ - MIRAFLORES",
						"REGIONAL LA PAZ - TEMBLADERANI",
						"REGIONAL LA PAZ - SAN PEDRO",
						"PAMPAHASI",
						"REGIONAL LA PAZ - PAMPAHASI",
						"BUENOS AIRES",
						"REGIONAL LA PAZ - BUENOS AIRES",
						"REGIONAL LA PAZ - COROICO",
						"REGIONAL LA PAZ - RURRENABAQUE",
						"REGIONAL LA PAZ - PERIFERICA",
						"REGIONAL LA PAZ - ZONA SUR",
						"REGIONAL LA PAZ - CAMACHO",
						"REGIONAL LA PAZ - SAN MIGUEL",
						"REGIONAL LA PAZ - COBIJA",
						"REGIONAL LA PAZ - CHULUMANI",
						"REGIONAL LA PAZ - CARANAVI"

					],
					
					"BFIE EL ALTO":[
						"REGIONAL EL ALTO - SENKATA",
						"REGIONAL EL ALTO - 21 DE OCTUBRE",
						"REGIONAL EL ALTO - ACHOCALLA",
						"REGIONAL EL ALTO - NUEVOS HORIZONTES",
						"REGIONAL EL ALTO - LIBERTAD",
						"REGIONAL EL ALTO - PATACAMAYA",
						"REGIONAL EL ALTO - TERMINAL",
						"REGIONAL EL ALTO - VILLA YUNGUYO",
						"REGIONAL EL ALTO - BALLIVIAN",
						"REGIONAL EL ALTO - DESAGUADERO",
						"REGIONAL EL ALTO - ACHACACHI",
						"REGIONAL EL ALTO - COPACABANA",
						"REGIONAL EL ALTO - VILLA ADELA",
						"REGIONAL EL ALTO - VILLA ESPERANZA",
						"REGIONAL EL ALTO - OFICINA CENTRAL",
						"REGIONAL EL ALTO - NUEVO AMANECER",
						"REGIONAL EL ALTO - PANAMERICANA",
						"REGIONAL EL ALTO - RIO SECO",
						"REGIONAL EL ALTO - LA CEJA",
						"REGIONAL EL ALTO - ROMERO PAMPA",
						"REGIONAL EL ALTO - VILLA DOLORES",
						"REGIONAL EL ALTO - VENTILLA",
						"REGIONAL EL ALTO - 16 DE JULIO",
						"REGIONAL EL ALTO - VENTILLA",
						"REGIONAL EL ALTO - VIACHA",
						"REGIONAL EL ALTO - 12 DE OCTUBRE",
						"REGIONAL EL ALTO - SATELITE",
						"REGIONAL EL ALTO - SANTIAGO",
						"REGIONAL EL ALTO - PACAJES"
					],

					"BFIE COCHABAMBA": [
						"REGIONAL COCHABAMBA - VINTO",
						"REGIONAL COCHABAMBA - TAMBORADA",
						"REGIONAL COCHABAMBA - CRUCE TAQUIÑA",
						"REGIONAL COCHABAMBA - QUILLACOLLO",
						"REGIONAL COCHABAMBA - SACABA",
						"REGIONAL COCHABAMBA - JORDAN",
						"REGIONAL COCHABAMBA - COLCAPIRHUA",
						"REGIONAL COCHABAMBA - PUNATA",
						"REGIONAL COCHABAMBA - PACATA",
						"REGIONAL COCHABAMBA - CALA CALA",
						"REGIONAL COCHABAMBA - KANATA",
						"REGIONAL COCHABAMBA - SAN MARTIN",
						"REGIONAL COCHABAMBA - LA CANCHA",
						"REGIONAL COCHABAMBA - NATANIEL AGUIRRE",
						"REGIONAL COCHABAMBA - HEROINAS",
						"REGIONAL COCHABAMBA - PETROLERA",
						"REGIONAL COCHABAMBA - ENTRE RIOS",
						"REGIONAL COCHABAMBA - ALALAY",
						"REGIONAL COCHABAMBA - IVIRGARZAMA",
						"REGIONAL COCHABAMBA - CHIMORE",
						"REGIONAL COCHABAMBA - URKUPIÑA",
						"REGIONAL COCHABAMBA - CRUCE TAQUIÑA"

					],

					"BFIE ORURO":[
						"REGIONAL ORURO - PAGADOR",
						"REGIONAL ORURO - SUCURSAL ORURO",
						"REGIONAL ORURO - ESPAÑA",
						"REGIONAL ORURO - LLALLAGUA",
						"REGIONAL ORURO - BOLIVAR"
					],

					"BFIE TARIJA":[
						"REGIONAL TARIJA - YACUIBA",
						"REGIONAL TARIJA - CAMARGO",
						"REGIONAL TARIJA - LA TABLADA",
						"REGIONAL TARIJA - ANDALUCIA",
						"REGIONAL TARIJA - LUIS DE FUENTE",
						"REGIONAL TARIJA - GUADALUPE",
						"REGIONAL TARIJA - BERMEJO",
						"REGIONAL TARIJA - GUADALQUIVIR",
						"REGIONAL TARIJA - VALLE DE CONCEPCION",
						"REGIONAL TARIJA - ARANJUEZ",
						"REGIONAL TARIJA - VILLAMONTES"

					],

					"BFIE SUCRE":[
						"REGIONAL SUCRE - PADILLA",
						"REGIONAL SUCRE - 25 DE MAYO",
						"REGIONAL SUCRE - CHARCAS",
						"REGIONAL SUCRE - MONTEAGUDO",
						"REGIONAL SUCRE - MESA VERDE"
					],

					"BFIE POTOSI":[
						"REGIONAL POTOSI - SAN ROQUE",
						"REGIONAL POTOSI - CERRO DE PLATA",
						"REGIONAL POTOSI - COTAGAITA",
						"REGIONAL POTOSI - TUPIZA",
						"REGIONAL POTOSI - 10 DE NOVIEMBRE",
						"REGIONAL POTOSI - UYUNI",
						"REGIONAL POTOSI - BETANZOS",
						"REGIONAL POTOSI - NUEVA TERMINAL",
						"REGIONAL POTOSI - JUNIN",
						"REGIONAL POTOSI - VILLAZON"

					],
					
            "BFIE SANTA CRUZ": [
                "124-2 - BFIE SANTA CRUZ"],
            "BFIE INTERNO": [   
                "C008 - POLYSISTEMAS"]
                        
        },

        "JTI BOLIVIA": {
            "JTI EXTERNO": [
                "367 - JTI BOLIVIA"],
            "JTI INTERNO": [   
                "C008 - POLYSISTEMAS"]
        },

        "REGIONAL COCHABAMBA": {
            "": [
                "301 - ESTEBAN ARCE",
                "302 - SAN MARTIN",
                "303 - HUAYRA KHASA",
                "305 - CRUCE TAQUIÑA",
                "306 - QUILLACOLLO",
                "307 - COLCA PIRHUA",
                "309 - MUYURINA",
                "310 - NORMALIZACIÓN COCHABAMBA",
                "311 - RECAUDADORA JORDAN",
                "312 - SACABA",
                "313 - VILLA GALINDO",
                "314 - PUNATA",
                "316 - AYACUCHO",
                "317 - SOL AMIGO COCHABAMBA",
                "318 - PANAMERICANA",
                "320 - CLIZA",
                "321 - VINTO",
                "322 - REPÚBLICA",
                "323 - TIQUIPAYA",
                "324 - QUINTANILLA",
                "325 - JORDAN",
                "326 - PLAZA BOLIVAR",
                "327 - PETROLERA",
                "328 - LA CHIMBA",
                "329 - AMERICA",
                "331 - EL AVION",
                "332 - VILLA PAGADOR",
                "333 - PACATA"
            ]
        },
        "REGIONAL BENI": {
            "": [
                "801 - TRINIDAD",
                "802 - RIBERALTA",
                "803 - GUAYARAMERIN"
            ]
        },
        "REGIONAL COBIJA": {
            "": [
                "901 - COBIJA",
                "902 - TAJIBOS",
                "903 - NORMALIZACIÓN"
            ]
        },
        "REGIONAL POTOSI": {
            "": [                                                
						"500 - REGIONAL POTOSI", 
                "501 - MERCADO UYUNI",
                "502 - SOL AMIGO POTOSI",
                "503 - BOULEVARD",
					    "505 - BANDERAS", 
					    "506 - MURILLO"
            ]
        },
        "REGIONAL SANTA CRUZ": {
            "": [
                "701 - CASCO VIEJO",
                "702 - EL PARI",
                "703 - MUTUALISTA",
                "704 - 1RO. DE MAYO",
                "705 - MONTERO",
                "706 - EL TORNO",
                "709 - PIRAI",
                "711 - PLAN 3000",
                "715 - LA GUARDIA",
                "716 - ALTO SAN PEDRO",
                "718 - NORTE",
                "719 - SOL AMIGO SANTA CRUZ",
                "721 - ARROYO CONCEPCIÓN",
                "723 - PAMPA DE LA ISLA",
                "724 - COLORADA",
                "725 - Minero",
                "726 - SAN JULIAN",
                "727 - SAN JOSE",
                "728 - 2 DE AGOSTO",
                "730 - LOS LOTES",
                "731 - GERMAN MORENO",
                "732 - YAPACANI",
                "734 - SATELITE NORTE",
                "735 - NORTE I",
                "736 - EL QUIOR",
                "737 - VIRGEN DE LUJAN",
                "738 - EL BAJIO",
                "739 - CRISTO REDENTOR"
            ]
        },
        "MOLINO ANDINO" :{
            "MOLINO ANDINO" : [
                "620 - MOLINO ANDINO S.A."]
	            },
        "ROCHE BOLIVIA" :{
            "ROCHE BOLIVIA" :[
                "599 - ROCHE BOLIVIA S.R.L."]
        }
    }

    if uploaded_file:
        df = pd.read_excel(uploaded_file)
                
        col1, col2 = st.columns(2)

        with col1:
            st.subheader("FILTRAR TIPO DE SOLICITUD")
            tipos_seleccionados = st.multiselect(
                "Selecciona uno o varios tipos de solicitud:",
                options=TIPOS_SOLICITUD,
                default=TIPOS_SOLICITUD
            )
        with col2:
            st.subheader("SELECCIONAR REGION Y SUBZONA")

            region = st.selectbox("Selecciona una región", options=list(CENTROS_CENTRALIZADO.keys()))
            subzonas = list(CENTROS_CENTRALIZADO[region].keys())
            subzona = st.selectbox("Selecciona una subzona", options=subzonas)

            centros_disponibles = CENTROS_CENTRALIZADO[region][subzona]
            centros_seleccionados = st.multiselect("Selecciona Centros de Costo", centros_disponibles, default=centros_disponibles)

        # Filtrado de datos
        df_filtrado = df[
            (df["Tipo de Solicitud"].isin(tipos_seleccionados)) &
            (df["Centro de Costo"].isin(centros_seleccionados))
        ]

        # Eliminar columnas no deseadas
        columnas_a_eliminar = ["Autorizado", "Locacion", "Centro de Costo Polysistemas", "Fecha de Impresión"]
        df_filtrado = df_filtrado.drop(columns=columnas_a_eliminar, errors='ignore')

        st.success(f"🔍 Filtrado completo: {len(df_filtrado)} registros encontrados.")

        colum1, colum2 = st.columns([15.50, 6.10])

        with colum1:
            # Mostrar tabla filtrada
            st.dataframe(df_filtrado, use_container_width=True)

            # Input para que el usuario elija el nombre del archivo
            nombre_archivo = st.text_input(
                "📝 Nombre del archivo a descargar:",
                value="resultado_filtrado.xlsx",
                placeholder="Ej: rutas_filtradas.xlsx"
            )

        with colum2:
            # Mostrar resumen basado en TipoFile y suma de Cantidad como tabla
            tipos_resumen = ["CAJA", "Caja", "Cintillos", "FILE"]
            st.subheader("📦 RESUMEN DE CANTIDADES")

            if "TipoFile" in df_filtrado.columns and "Cantidad" in df_filtrado.columns:
                resumen_data = []

                for tipo in tipos_resumen:
                    total = df_filtrado[df_filtrado["TipoFile"] == tipo]["Cantidad"].sum()
                    resumen_data.append({"TipoFile": tipo, "Total Cantidad": int(total)})

                resumen_df = pd.DataFrame(resumen_data)
                st.table(resumen_df)
            else:
                st.warning("⚠️ Las columnas 'TipoFile' o 'Cantidad' no se encuentran en el archivo.")
                    
            # -------------------------------
            # BOTÓN PARA DESCARGAR RESUMEN
            # -------------------------------
            resumen_output = BytesIO()
            with pd.ExcelWriter(resumen_output, engine='openpyxl') as writer:
                resumen_df.to_excel(writer, index=False, sheet_name='Resumen')
            resumen_output.seek(0)

            st.download_button(
                label="📥 Descargar Resumen de Cantidades",
                data=resumen_output,
                file_name="resumen_cantidades.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

            #else:
            #    st.warning("⚠️ Las columnas 'TipoFile' o 'Cantidad' no se encuentran en el archivo.")


        # Descarga directa sin guardar en carpeta
        st.subheader("📥 Descargar Excel Filtrado")

        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_filtrado.to_excel(writer, index=False, sheet_name='Filtrado')
        output.seek(0)

        st.download_button(
            label="📥 Descargar Excel filtrado",
            data=output,
            #file_name="resultado_filtrado.xlsx",
            file_name=nombre_archivo if nombre_archivo.strip() else "resultado_filtrado.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )


# --------------------- 2. COMBINAR ARCHIVOS ---------------------
def combinar_archivos():
    st.title("📚 COMBINAR ARCHIVOS FILTRADOS")
              

    co1,co2 = st.columns([15.04, 8.05])
    with co1:
        # Cargar múltiples archivos
        uploaded_files = st.file_uploader("Elige archivos Excel", type=["xlsx", "xls"], accept_multiple_files=True)
                
    #with co2:
        # Input para ruta personalizada (opcional)
        #carpeta_guardado = st.text_input(
            #"📁 Ingresa la ruta para para guardar el Archivo Filtrado :",
            #placeholder="Ej: C:/Users/TuUsuario/Desktop"
        #)

    # Comprobar si se han subido archivos
    if uploaded_files:

        # Ordenar los archivos por nombre, si es necesario
        uploaded_files = sorted(uploaded_files, key=lambda x: x.name)

        dfs = []
        for file in uploaded_files:
            # Leer cada archivo Excel en un DataFrame
            df = pd.read_excel(file)
            dfs.append(df)

        # Combinar todos los DataFrames en uno solo
        combined_df = pd.concat(dfs, ignore_index=True)

        # Mostrar el DataFrame combinado
        st.write("DataFrame Combinado:")
        st.dataframe(combined_df)

        # Función para convertir el DataFrame combinado a Excel
        def convert_df_to_excel(df):
            # Crear un objeto BytesIO
            output = BytesIO()
            # Escribir el DataFrame en el objeto BytesIO
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False)
            # Mover el cursor al principio del objeto BytesIO
            output.seek(0)
            return output
                
        excel_data = convert_df_to_excel(combined_df)

        # Botón de descarga
        st.download_button(
            label="📥 Descargar Excel Combinado",
            data=excel_data,
            file_name="Archivo_Combinado.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
                
    else:
        st.write("Por favor, sube los archivos Excel para combinarlos.")


# --------------------- 3. GENERADOR DE SINTAXIS ---------------------
def generar_sintaxis():

    st.title("🔎 GENERADOR DE SINTAXIS DE BÚSQUEDA")

    col1, col2 = st.columns((2))
    with col1:
        # Cargar archivo Excel
        st.markdown("<h4>Selecciona el archivo Excel con la columna 'SolicitudCode'</h4>", unsafe_allow_html=True)
        uploaded_file = st.file_uploader("Selecciona un archivo .xlsx", type="xlsx")

        # Variable para almacenar la sintaxis generada
        sintaxis_final = ""

        # Generar la sintaxis de búsqueda
        if st.button("Generar Sintaxis"):

            if uploaded_file is not None:
                try:
                    # Leer el archivo Excel y obtener solo la columna 'SolicitudCode'
                    df = pd.read_excel(uploaded_file)

                    # Verificar si la columna existe
                    if 'SolicitudCode' not in df.columns:
                        st.error("La columna 'SolicitudCode' no se encuentra en el archivo.")
                    else:
                        # Eliminar las demás columnas y eliminar nulos
                        solicitud_codes = df['SolicitudCode'].dropna().astype(str).tolist()

                        # Crear la sintaxis
                        solicitud_sintaxis = " | ".join(
                            [f'\n{{[Solicitud]:[Nro Solicitud]="{code.strip()}"}}' for code in solicitud_codes if code.strip()]
                        )

                        sintaxis_final = f"{solicitud_sintaxis}"
                except Exception as e:
                    st.error(f"Error al leer el archivo: {e}")

    with col2:
        # Mostrar la sintaxis generada
        if sintaxis_final:
            st.success("Sintaxis de búsqueda generada:")
            st.code(sintaxis_final)

            # Convertir a bytes para evitar errores de descarga
            contenido_descarga = sintaxis_final.encode("utf-8")

            st.download_button(
                label="📥 Descargar sintaxis",
                data=contenido_descarga,
                file_name="sintaxis_busqueda.txt",
                mime="text/plain"
            )

# --------------------- 4. PREPARAR RUTA ALMACÉN V2 ---------------------
def preparar_ruta():

    import streamlit as st
    import pandas as pd
    from io import BytesIO
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    st.title("📦 PREPARAR RUTA ALMACÉN LA PAZ")

    col1, col2 = st.columns(2)
    file1 = col1.file_uploader("📂 Cargar Archivo 1", type="xlsx")
    file2 = col2.file_uploader("📂 Cargar Archivo 2", type="xlsx")

    with col1:
        file_name_input = st.text_input("📝 Nombre del archivo para CORREO", value="Archivo_Enviar_CORREO")
        file_name_correo = st.text_input("📝 Nombre archivo para RUTA", value="Archivo_Enviar_Ruta")
        sheet_name_input = st.text_input("📄 Nombre de la hoja", value="Ruta")

    with col2:

        recojo_title = st.text_input("✏️ Título Recojo", value="RECOJO DE CAJAS B_SOL")
        devolucion_title = st.text_input("✏️ Título Devolución", value="DEVOLUCIÓN DE ITEMS B_SOL")

    if file1 and file2:

        df1 = pd.read_excel(file1)
        df1.drop(['Estado de Solicitud', 'Servicio', 'Turno', 'Tipo de Recojo', 'Fecha de Registro'], 
                axis=1, inplace=True, errors='ignore')

        df2 = pd.read_excel(file2)

        df1['SolicitudCode'] = df1['SolicitudCode'].astype(str).str.strip()
        df2['Nro Solicitud'] = df2['Nro Solicitud'].astype(str).str.strip()

        merged_df = pd.merge(
            df1,
            df2[['Nro Solicitud', 'Solicitante', 'Usuario', 'Items Oneil']],
            left_on='SolicitudCode',
            right_on='Nro Solicitud',
            how='left'
        )

        if 'Centro de Costo' in merged_df.columns:
            merged_df = merged_df.sort_values(by='Centro de Costo')

        df_recojo = merged_df[merged_df['TipoFile'] != 'FILE'].copy()
        df_devolucion = merged_df[merged_df['TipoFile'] == 'FILE'].copy()
                
        # 🔍 MOSTRAR PRIMERAS 5 FILAS
        st.subheader("🔍 VISTA PREVIA PARA PREPARAR GUIA")
        st.dataframe(merged_df.head())

        st.subheader("📦 – RECOJO")
        st.dataframe(df_recojo.head())

        st.subheader("📦 – DEVOLUCIÓN")
        st.dataframe(df_devolucion.head())

        for df in [df_recojo, df_devolucion, merged_df]:
            df.drop(['Usuario', 'Nro Solicitud'], axis=1, inplace=True, errors='ignore')
            df.dropna(how="all", inplace=True)
            df.replace("", pd.NA, inplace=True)
            df.dropna(how="all", inplace=True)

        # ======================================================
        # FUNCIÓN BORDES
        # ======================================================
        thin_border = Border(
            left=Side(border_style="thin", color="000000"),
            right=Side(border_style="thin", color="000000"),
            top=Side(border_style="thin", color="000000"),
            bottom=Side(border_style="thin", color="000000")
        )

        def aplicar_bordes(ws, start_row, df):
            nrows, ncols = df.shape
            for row in range(start_row, start_row + nrows + 1):
                for col in range(1, ncols + 1):
                    ws.cell(row=row, column=col).border = thin_border

        # ======================================================
        # FUNCIÓN AUTOAJUSTE COLUMNAS
        # ======================================================
        def auto_ajustar_columnas(ws):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = max_len + 3

        # ======================================================
        # ARCHIVO CORREO
        # ======================================================
        def generar_excel_principal(df_recojo, df_devolucion):
            output = BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                workbook = writer.book
                ws = workbook.create_sheet(sheet_name_input)

                if "Sheet" in workbook.sheetnames:
                    del workbook["Sheet"]

                color_lila = "C27BA0"
                color_rosa = "F4B1BA"
                bold_font = Font(bold=True)

                current_row = 1

                # ------------------ RECOJO ------------------
                last_col = get_column_letter(len(df_recojo.columns))
                ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
                cell = ws[f"A{current_row}"]
                cell.value = recojo_title
                cell.font = Font(bold=True, size=14)
                cell.alignment = Alignment(horizontal="center")
                cell.fill = PatternFill("solid", fgColor=color_lila)

                startrow = current_row + 2
                df_recojo.to_excel(writer, sheet_name=sheet_name_input, startrow=startrow, index=False)
                header_row = startrow + 1

                for col in range(1, len(df_recojo.columns) + 1):
                    h = ws.cell(row=header_row, column=col)
                    h.fill = PatternFill("solid", fgColor=color_lila)
                    h.font = bold_font

                aplicar_bordes(ws, header_row, df_recojo)

                current_row = header_row + len(df_recojo) + 3

                # ------------------ DEVOLUCIÓN ------------------
                last_col = get_column_letter(len(df_devolucion.columns))
                ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
                cell = ws[f"A{current_row}"]
                cell.value = devolucion_title
                cell.font = Font(bold=True, size=14)
                cell.alignment = Alignment(horizontal="center")
                cell.fill = PatternFill("solid", fgColor=color_rosa)

                startrow = current_row + 2
                df_devolucion.to_excel(writer, sheet_name=sheet_name_input, startrow=startrow, index=False)
                header_row = startrow + 1

                for col in range(1, len(df_devolucion.columns) + 1):
                    h = ws.cell(row=header_row, column=col)
                    h.fill = PatternFill("solid", fgColor=color_rosa)
                    h.font = bold_font

                aplicar_bordes(ws, header_row, df_devolucion)

                # 🔥 Ajuste automático de columnas para todo el libro
                auto_ajustar_columnas(ws)

            output.seek(0)
            return output

        # ======================================================
        # ARCHIVO RUTA
        # ======================================================
        def generar_excel_correo(df):
            columnas = [
                'SolicitudCode',
                'Tipo de Solicitud',
                'Cliente',
                'WorkOrderCode',
                'Cantidad',
                'TipoFile',
                'Centro de Costo',
                'Solicitante',
                'Items Oneil'
            ]

            df_simple = df[columnas].copy()

            output = BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df_simple.to_excel(writer, index=False, sheet_name="RUTA")
                ws = writer.book["RUTA"]

                aplicar_bordes(ws, 1, df_simple)
                auto_ajustar_columnas(ws)   # 🔥 Autoajuste aquí también

            output.seek(0)
            return output

        # ======================================================
        # DESCARGAS
        # ======================================================
        archivo_principal = generar_excel_principal(df_recojo, df_devolucion)
        archivo_correo = generar_excel_correo(merged_df)

        st.download_button(
            "📥 Descargar Archivo para CORREO",
            archivo_principal,
            file_name=file_name_input + ".xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        st.download_button(
            "📧 Descargar Archivo para RUTA",
            archivo_correo,
            file_name=file_name_correo + ".xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )


# --------------------- 5. PREPARAR GUIA PARA RUTA V3 ---------------------
def preparar_guia_ruta():
    import streamlit as st
    import pandas as pd
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.worksheet.page import PageMargins
    import datetime
    import re

    st.title("📦 GENERADOR DE GUÍA - FORMATO OFICIAL POLYSISTEMAS")
    st.caption("Genera una guía ordenada para imprimir, guardar como PDF desde Excel o adjuntar a un correo.")

    col1, col2 = st.columns(2)
    with col1:
        logo_file = st.file_uploader("📸 Subir logo (PNG o JPG)", type=["png", "jpg", "jpeg"], key="guia_logo")
    with col2:
        uploaded_file = st.file_uploader("📄 Cargar archivo Excel para generar la ruta", type=["xlsx"], key="guia_ruta_excel")

    if not uploaded_file:
        st.info("Carga el archivo Excel para configurar y generar la guía.")
        return

    try:
        df = pd.read_excel(uploaded_file)
    except Exception as e:
        st.error(f"No se pudo leer el archivo Excel: {e}")
        return

    df.columns = df.columns.str.strip().str.replace(" ", "_").str.replace("-", "_")
    required_columns = ["Solicitante", "Centro_de_Costo", "WorkOrderCode", "Cantidad", "Items_Oneil", "TipoFile"]
    missing = [col for col in required_columns if col not in df.columns]
    if missing:
        st.error(f"Faltan columnas requeridas: {', '.join(missing)}")
        return

    df["Items_Oneil"] = df["Items_Oneil"].fillna("CINTILLOS")
    df["Cantidad"] = pd.to_numeric(df["Cantidad"], errors="coerce").fillna(0)

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        personal = st.selectbox("👤 Personal de Polysistemas", [
            "JAIME QUISPE", "CARLOS ORTIZ", "MARCO HUAYLLUCO", "ALFREDO RIVEROS"
        ])
    with c2:
        cliente = st.selectbox("🏢 Cliente", ["BANCO SOL", "BNB", "BANCO FIE"])
    with c3:
        tipo_seleccionado = st.multiselect(
            "📦 Tipos de archivo a procesar",
            ["CAJA", "Caja", "Cintillos", "FILE"],
            default=["CAJA"]
        )
    with c4:
        regional = st.selectbox("🌎 Regional", ["LA PAZ", "EL ALTO", "ORURO", "POTOSI", "SUCRE", "VIACHA"])

    fecha_actual = datetime.datetime.now().strftime("%d/%m/%Y")
    df = df[df["TipoFile"].isin(tipo_seleccionado)].copy()

    if df.empty:
        st.warning("No hay registros para los tipos de archivo seleccionados.")
        return

    st.info(
        f"**Personal:** {personal}  \n"
        f"**Cliente:** {cliente}  \n"
        f"**Fecha:** {fecha_actual}  \n"
        f"**Regional:** {regional}  \n"
        f"**Tipos seleccionados:** {', '.join(tipo_seleccionado)}"
    )
    st.subheader("🔍 Vista previa")
    st.dataframe(df.head(10), use_container_width=True)

    def obtener_rango_codigos(cadena, tipo_file):
        if tipo_file != "Caja":
            if pd.isna(cadena) or str(cadena).strip().upper() in ("", "CINTILLOS", "NONE", "NAN"):
                return "CINTILLOS"
            return str(cadena).strip()
        if pd.isna(cadena) or str(cadena).strip().upper() in ("", "CINTILLOS", "NONE", "NAN"):
            return "CINTILLOS"
        partes = [p.strip() for p in re.split(r"[;,]", str(cadena)) if p.strip().isdigit()]
        if not partes:
            return "CINTILLOS"
        partes = sorted(int(p) for p in partes)
        return f"{partes[0]} - {partes[-1]}"

    # Crear libro de Excel con diseño profesional y configuración de impresión.
    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Guía"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.45, bottom=0.45, header=0.2, footer=0.2)
    ws.oddFooter.center.text = "Polysistemas | Guía de recepción de mensajería"
    ws.oddFooter.right.text = "Página &P de &N"
    ws.print_options.horizontalCentered = True

    rojo_hex = "933C47"
    rojo = PatternFill("solid", fgColor=rojo_hex)
    gris = PatternFill("solid", fgColor="F2F2F2")
    gris_oscuro = PatternFill("solid", fgColor="D9E1F2")
    blanco_bold = Font(name="Arial", size=9, color="FFFFFF", bold=True)
    bold = Font(name="Arial", size=9, bold=True)
    normal = Font(name="Arial", size=9)
    titulo = Font(name="Arial", size=13, bold=True, color="262626")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="777777")
    medium = Side(style="medium", color=rojo_hex)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def aplicar_bordes_fila(row, cols=7):
        for c in range(1, cols + 1):
            ws.cell(row, c).border = border

    def estilo_celda(cell, font=None, fill=None, alignment=None):
        cell.font = font or normal
        cell.border = border
        if fill:
            cell.fill = fill
        cell.alignment = alignment or center

    if logo_file:
        try:
            img = XLImage(logo_file)
            img.width = 175
            img.height = 62
            ws.add_image(img, "A1")
        except Exception:
            st.warning("No se pudo insertar el logo. La guía se generará sin logo.")

    # Encabezado institucional
    ws.merge_cells("C1:E1")
    ws["C1"] = "FORMATO"
    ws["C1"].font = bold
    ws["C1"].alignment = center
    ws.merge_cells("C2:E3")
    ws["C2"] = "GUÍA DE RECEPCIÓN DE MENSAJERÍA"
    ws["C2"].font = titulo
    ws["C2"].alignment = center
    for row in range(1, 4):
        for col in range(3, 6):
            ws.cell(row, col).border = border

    meta = {
        "F1": "Código:", "G1": "BOL-2.18-FOR-01",
        "F2": "Edición:", "G2": "02",
        "F3": "Fecha edición:", "G3": "26/11/2025"
    }
    for cell_ref, value in meta.items():
        ws[cell_ref] = value
        estilo_celda(ws[cell_ref], bold if cell_ref.startswith("F") else normal,
                     gris if cell_ref.startswith("F") else None, center)

    # Personal y cliente
    for row, label, value, right_label, right_value in [
        (5, "PERSONAL DE POLYSISTEMAS", personal, "FECHA:", fecha_actual),
        (6, "CLIENTE:", cliente, "CANTIDAD TOTAL:", float(df["Cantidad"].sum()))
    ]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row, 1, label)
        estilo_celda(ws.cell(row, 1), blanco_bold, rojo, left_wrap)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
        ws.cell(row, 4, value)
        estilo_celda(ws.cell(row, 4), bold, None, center)
        ws.cell(row, 6, right_label)
        estilo_celda(ws.cell(row, 6), blanco_bold, rojo, center)
        ws.cell(row, 7, right_value)
        estilo_celda(ws.cell(row, 7), bold, None, center)
        ws.row_dimensions[row].height = 24

    ws.merge_cells("A8:B8")
    ws["A8"] = "REGIONAL:"
    estilo_celda(ws["A8"], bold, gris, left_wrap)
    ws.merge_cells("C8:E8")
    ws["C8"] = regional
    estilo_celda(ws["C8"], bold, None, center)
    ws.merge_cells("F8:G8")
    ws["F8"] = "DOCUMENTO DE CONTROL DE ENTREGA"
    estilo_celda(ws["F8"], bold, gris_oscuro, center)
    aplicar_bordes_fila(8)

    start_row = 9
    solicitantes = df["Solicitante"].dropna().unique().tolist()
    if not solicitantes:
        solicitantes = ["SIN SOLICITANTE"]

    for solicitante in solicitantes:
        df_s = df[df["Solicitante"] == solicitante].copy()
        if df_s.empty and solicitante != "SIN SOLICITANTE":
            continue
        agencia = str(df_s["Centro_de_Costo"].iloc[0]) if not df_s.empty else ""

        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2)
        ws.cell(start_row, 1, "SOLICITANTE:")
        estilo_celda(ws.cell(start_row, 1), bold, gris, left_wrap)
        ws.merge_cells(start_row=start_row, start_column=3, end_row=start_row, end_column=5)
        ws.cell(start_row, 3, str(solicitante))
        estilo_celda(ws.cell(start_row, 3), bold, None, center)
        ws.cell(start_row, 6, "AGENCIA:")
        estilo_celda(ws.cell(start_row, 6), bold, gris, center)
        ws.cell(start_row, 7, agencia)
        estilo_celda(ws.cell(start_row, 7), bold, None, center)
        ws.row_dimensions[start_row].height = 25
        aplicar_bordes_fila(start_row)
        start_row += 1

        headers = ["NRO", "NRO. WO", "CANTIDAD", "CONTAINER CODE", "OBSERVACIONES", "FIRMA ENTREGA", "FIRMA RECEPCIÓN"]
        for col, header in enumerate(headers, 1):
            ws.cell(start_row, col, header)
            estilo_celda(ws.cell(start_row, col), blanco_bold, rojo, center)
        ws.row_dimensions[start_row].height = 30
        start_row += 1

        total_solic = 0
        for nro, r in enumerate(df_s.itertuples(index=False), start=1):
            vals = [
                nro,
                getattr(r, "WorkOrderCode", ""),
                getattr(r, "Cantidad", 0),
                obtener_rango_codigos(getattr(r, "Items_Oneil", "CINTILLOS"), getattr(r, "TipoFile", "")),
                "",
                "",
                ""
            ]
            for col, value in enumerate(vals, 1):
                cell = ws.cell(start_row, col, value)
                estilo_celda(cell, normal, None, center if col in (1, 2, 3, 6, 7) else left_wrap)
            ws.row_dimensions[start_row].height = 29
            total_solic += float(getattr(r, "Cantidad", 0) or 0)
            start_row += 1

        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2)
        ws.cell(start_row, 1, "TOTAL")
        estilo_celda(ws.cell(start_row, 1), bold, gris_oscuro, center)
        ws.cell(start_row, 3, total_solic)
        estilo_celda(ws.cell(start_row, 3), bold, gris_oscuro, center)
        ws.merge_cells(start_row=start_row, start_column=4, end_row=start_row, end_column=5)
        ws.cell(start_row, 4, f"REGISTROS: {len(df_s)}")
        estilo_celda(ws.cell(start_row, 4), bold, gris_oscuro, center)
        ws.cell(start_row, 6, "ENTREGUÉ CONFORME")
        estilo_celda(ws.cell(start_row, 6), bold, gris, center)
        ws.cell(start_row, 7, "RECIBÍ CONFORME")
        estilo_celda(ws.cell(start_row, 7), bold, gris, center)
        aplicar_bordes_fila(start_row)
        ws.row_dimensions[start_row].height = 32
        start_row += 2

    # Dimensiones optimizadas para lectura e impresión
    widths = [8, 18, 12, 34, 25, 20, 20]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 24
    ws.freeze_panes = "A10"
    ws.print_area = f"A1:G{max(start_row - 1, 8)}"
    ws.sheet_properties.outlinePr.summaryBelow = True

    wb.save(output)
    output.seek(0)
    safe_date = datetime.datetime.now().strftime("%Y%m%d")
    st.success("✅ Guía preparada con formato profesional y configuración para impresión.")
    st.download_button(
        label="📥 Descargar Guía Profesional (Excel)",
        data=output.getvalue(),
        file_name=f"Guia_Polysistemas_{safe_date}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )
    st.caption("Para enviar por correo en PDF: abre el Excel descargado y selecciona Archivo → Exportar/Guardar como PDF. La hoja está configurada en horizontal y ajustada al ancho de página.")

# --------------------- 6. ELIMINAR ARCHIVOS ---------------------
def eliminar_archivos():
    import streamlit as st
    import os
    import glob
    import pandas as pd

    #st.set_page_config(page_title="Eliminar Excel con vista previa", layout="centered")
    st.title("🗑️ Eliminar archivos Excel con vista previa desplegable")

    folder_input = st.text_input("Introduce la ruta de la carpeta con archivos Excel:")
    if st.button("🔄 Cargar archivos"):
        st.session_state.folder = folder_input
        st.session_state.selected_files = []

    # Inicializar estados
    if 'folder' not in st.session_state:
        st.session_state.folder = ""
    if 'selected_files' not in st.session_state:
        st.session_state.selected_files = []

    folder = st.session_state.folder

    def toggle_selection(f_name):
        sel = st.session_state.selected_files
        if f_name in sel:
            sel.remove(f_name)
        else:
            sel.append(f_name)

    if folder:
        if os.path.isdir(folder):
            files = glob.glob(os.path.join(folder, "*.xlsx")) + glob.glob(os.path.join(folder, "*.xls"))
            if files:
                st.subheader("📂 Selecciona archivos:")
                for path in files:
                    name = os.path.basename(path)
                    selected = name in st.session_state.selected_files
                    col1, col2 = st.columns([4,1])
                    with col1:
                        st.write(name)
                    with col2:
                        if st.button("✅" if selected else "☐", key=name):
                            toggle_selection(name)
                            st.rerun()

                if st.session_state.selected_files:
                    st.subheader("👁️ Previsualizaciones desplegables")
                    for name in st.session_state.selected_files:
                        file_path = os.path.join(folder, name)
                        with st.expander(f"Ver preview: {name}", expanded=False):
                            try:
                                df = pd.read_excel(file_path, nrows=10)
                                st.dataframe(df)
                            except Exception as e:
                                st.error(f"Error leyendo {name}: {e}")

                    if st.checkbox("⚠️ Confirmo que deseo eliminar los archivos seleccionados"):
                        if st.button("🗑️ Eliminar archivos"):
                            cnt = 0
                            for name in list(st.session_state.selected_files):
                                p = os.path.join(folder, name)
                                if os.path.exists(p):
                                    os.remove(p)
                                    cnt += 1
                            st.success(f"Se eliminaron {cnt} archivos correctamente.")
                            st.session_state.selected_files = []
                            st.session_state.folder = ""
                            st.rerun()
                else:
                    st.info("Selecciona archivos pulsando en los botones de la derecha.")
            else:
                st.warning("No se encontraron archivos Excel en la carpeta.")
        else:
            st.error("La ruta no es válida o no corresponde a una carpeta.")
    else:
        st.info("Introduce la ruta y pulsa 'Cargar archivos' para empezar.")

                                
# --------------------- EJECUTAR SEGÚN MENÚ ---------------------
if menu == "🌎 Detalle General":
    detalle_general()
elif menu == "🏷️ Filtrar Solicitudes":
    filtrar_solicitudes()
elif menu == "🧩 Combinar Archivos":
    combinar_archivos()
elif menu == "🔎 Generar Sintaxis de Búsqueda":
    generar_sintaxis()
elif menu == "📦 Preparar Ruta Almacén La Paz":
    preparar_ruta()
elif menu == "📦 Preparar Guias de Ruta":
    preparar_guia_ruta()
elif menu == "🗑️ Eliminar Archivos":
    eliminar_archivos()
