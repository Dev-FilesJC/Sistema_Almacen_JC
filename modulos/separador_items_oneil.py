# Se ejecuta dentro del espacio global de app.py para conservar compatibilidad.
import re
import hashlib
from io import BytesIO
import pandas as pd
import streamlit as st
from openpyxl.styles import PatternFill, Font

st.subheader("📦 SEPARADOR DE ITEMS ONEIL")
st.caption("Separa los códigos de LASER-BNB y los compara con Documentos Prestados. También puede incorporar el Estado de OT desde un tercer Excel.")

st.markdown("""
<style>
div[data-testid="stDownloadButton"] button, div[data-testid="stButton"] button {
 min-height:34px!important;height:auto!important;padding:.35rem .75rem!important;
 font-size:.88rem!important;white-space:normal!important;
}
</style>
""", unsafe_allow_html=True)


def separar_codigos(valor):
    if valor is None:
        return []
    try:
        if pd.isna(valor):
            return []
    except (TypeError, ValueError):
        pass
    texto = str(valor).strip()
    if not texto:
        return []
    return [parte.strip() for parte in re.split(r"[,;\n\r]+", texto) if parte.strip()]


def normalizar_codigo(valor):
    if valor is None:
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    texto = str(valor).strip()
    # Evita que Excel convierta un código numérico en '123.0'.
    if re.fullmatch(r"\d+\.0", texto):
        texto = texto[:-2]
    return re.sub(r"\s+", "", texto).upper()


def normalizar_columna(valor):
    return re.sub(r"\s+", " ", str(valor).replace("\n", " ").strip()).casefold()


def buscar_columna(df, nombres):
    mapa = {normalizar_columna(c): c for c in df.columns}
    for nombre in nombres:
        if normalizar_columna(nombre) in mapa:
            return mapa[normalizar_columna(nombre)]
    return None


def leer_excel(upload):
    upload.seek(0)
    frame = pd.read_excel(upload, dtype=str, keep_default_na=False)
    frame.columns = [str(col).strip() for col in frame.columns]
    return frame


def crear_excel(dataframe, sheet_name="Resultado", columna_marca=None):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        dataframe.to_excel(writer, index=False, sheet_name=sheet_name[:31])
        ws = writer.sheets[sheet_name[:31]]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        encabezados = {str(c.value).strip(): c.column for c in ws[1]}
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.value is not None:
                    cell.number_format = "@"
        # Marca visualmente todas las filas duplicadas/encontradas.
        if columna_marca and columna_marca in encabezados:
            idx = encabezados[columna_marca]
            for fila in range(2, ws.max_row + 1):
                marca = str(ws.cell(fila, idx).value or "").strip().upper()
                if marca in {"DUPLICADO", "ENCONTRADO", "SÍ", "SI"}:
                    for celda in ws[fila]:
                        celda.fill = PatternFill("solid", fgColor="FFF2CC")
                        celda.font = Font(color="9C0006" if marca == "DUPLICADO" else "7F6000")
        for column_cells in ws.columns:
            letter = column_cells[0].column_letter
            max_len = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)
    output.seek(0)
    return output.getvalue()


# 1. Archivo principal: LASER-BNB (o cualquier Excel con las columnas equivalentes)
uploaded_file = st.file_uploader(
    "1️⃣ Sube el Excel principal (ej. LASER-BNB.xlsx)",
    type=["xlsx", "xls"], key="separador_items_archivo",
    help="Se necesitan las columnas Nombre, Fecha de creación e Items Oneil."
)

if uploaded_file is None:
    for key in ("resultado_items", "items_source_id", "items_result_source_id", "resultado_comparacion", "comparacion_source_id"):
        st.session_state.pop(key, None)
    st.info("Carga el Excel principal para comenzar.")
else:
    try:
        source_id = hashlib.sha256(uploaded_file.getvalue()).hexdigest()
        if st.session_state.get("items_source_id") != source_id:
            for key in ("resultado_items", "sin_codigo_items", "nombre_col_items", "items_col_items",
                        "items_result_source_id", "resultado_comparacion", "comparacion_source_id"):
                st.session_state.pop(key, None)
            st.session_state["items_source_id"] = source_id

        df = leer_excel(uploaded_file)
        st.success(f"Archivo principal cargado: {uploaded_file.name}")
        st.caption(f"Filas originales: {len(df):,} · Columnas: {len(df.columns)}")
        if df.empty:
            st.warning("El archivo no contiene filas de datos.")
        else:
            with st.expander("👁️ Vista previa del archivo principal", expanded=False):
                st.dataframe(df.head(10), use_container_width=True, hide_index=True)

            cols = df.columns.tolist()
            c1, c2, c3 = st.columns(3)
            with c1:
                nombre_default = buscar_columna(df, ["Nombre"])
                nombre_col = st.selectbox("Columna Nombre", cols, index=cols.index(nombre_default) if nombre_default else 0, key="items_nombre_col")
            with c2:
                fecha_default = buscar_columna(df, ["Fecha de creación", "Fecha de creacion"])
                fecha_col = st.selectbox("Columna Fecha de creación", cols, index=cols.index(fecha_default) if fecha_default else 0, key="items_fecha_col")
            with c3:
                items_default = buscar_columna(df, ["Items Oneil", "Items Oneil "])
                items_col = st.selectbox("Columna Items Oneil", cols, index=cols.index(items_default) if items_default else min(1, len(cols)-1), key="items_codigo_col")

            if st.button("✂️ Separar códigos", type="primary", key="btn_separar_items"):
                filas, sin_codigo = [], 0
                for _, row in df.iterrows():
                    codigos = separar_codigos(row[items_col])
                    if not codigos:
                        sin_codigo += 1
                        continue
                    for codigo in codigos:
                        nueva = row.to_dict()
                        nueva[nombre_col] = str(row[nombre_col]).strip()
                        nueva[fecha_col] = str(row[fecha_col]).strip()
                        nueva[items_col] = codigo
                        filas.append(nueva)
                st.session_state["resultado_items"] = pd.DataFrame(filas, columns=df.columns)
                st.session_state["sin_codigo_items"] = sin_codigo
                st.session_state["nombre_col_items"] = nombre_col
                st.session_state["fecha_col_items"] = fecha_col
                st.session_state["items_col_items"] = items_col
                st.session_state["items_result_source_id"] = source_id
                st.session_state.pop("resultado_comparacion", None)
                st.session_state.pop("comparacion_source_id", None)

            if st.session_state.get("items_result_source_id") == source_id and "resultado_items" in st.session_state:
                resultado = st.session_state["resultado_items"].copy()
                nombre_resultado = st.session_state["nombre_col_items"]
                fecha_resultado = st.session_state["fecha_col_items"]
                items_resultado = st.session_state["items_col_items"]
                st.divider()
                st.markdown("### 📋 Resultado de la separación")
                m1, m2, m3 = st.columns(3)
                m1.metric("Filas originales", f"{len(df):,}")
                m2.metric("Códigos separados", f"{len(resultado):,}")
                m3.metric("Filas sin código", f"{st.session_state['sin_codigo_items']:,}")
                if resultado.empty:
                    st.warning("No se encontraron códigos para separar.")
                else:
                    st.dataframe(resultado, use_container_width=True, hide_index=True)
                    reducido = resultado[[nombre_resultado, fecha_resultado, items_resultado]].copy()
                    reducido.columns = ["Nombre", "Fecha de creación", "Items Oneil"]
                    b1, b2 = st.columns(2)
                    with b1:
                        st.download_button("📥 Descargar Excel completo", data=crear_excel(resultado, "Items separados"),
                            file_name="LASERFICHE_ITEMS_SEPARADOS_COMPLETO.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="descargar_items_completo")
                    with b2:
                        st.download_button("📥 Descargar Nombre, Fecha e Items", data=crear_excel(reducido, "Items separados"),
                            file_name="LASERFICHE_NOMBRE_FECHA_ITEMS.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="descargar_items_reducido")

                    st.divider()
                    st.markdown("### 🔎 Comparar con Documentos Prestados")
                    st.write("Sube el archivo de préstamos. Se comparará **Items Oneil** con **Codigo del Item**. Los códigos presentes en ambos archivos se marcarán como **DUPLICADO** y se resaltarán en amarillo en la vista previa y en Excel.")
                    compare_file = st.file_uploader("2️⃣ Sube Documentos Prestados.xlsx", type=["xlsx", "xls"], key="comparar_items_archivo")
                    status_file = st.file_uploader("3️⃣ (Opcional) Sube el archivo de estado de Laserfiche (ej. stado-laser.xlsx)", type=["xlsx", "xls"], key="estado_laser_archivo",
                        help="Se usará la columna Estado de OT. Se relaciona por Nro Solicitud y, si no existe, por Nombre.")

                    if compare_file is not None:
                        compare_id = hashlib.sha256(compare_file.getvalue()).hexdigest()
                        try:
                            df_compare = leer_excel(compare_file)
                            col_codigo = buscar_columna(df_compare, ["Codigo del Item", "Código del Item"])
                            col_solicitud = buscar_columna(df_compare, ["Nro de Solicitud", "Nro Solicitud", "Nro. Solicitud"])
                            col_fecha_solicitud = buscar_columna(df_compare, ["Fecha de Solicitud", "Fecha d Solicitud"])
                            faltantes = []
                            if not col_codigo: faltantes.append("Codigo del Item")
                            if not col_solicitud: faltantes.append("Nro de Solicitud")
                            if not col_fecha_solicitud: faltantes.append("Fecha de Solicitud")
                            if faltantes:
                                st.error("En Documentos Prestados faltan estas columnas: " + ", ".join(faltantes))
                            else:
                                with st.expander("👁️ Vista previa de Documentos Prestados", expanded=False):
                                    st.dataframe(df_compare.head(10), use_container_width=True, hide_index=True)
                                if status_file is not None:
                                    df_status = leer_excel(status_file)
                                    status_id = hashlib.sha256(status_file.getvalue()).hexdigest()
                                    col_estado = buscar_columna(df_status, ["Estado de OT", "Estado OT"])
                                    status_solicitud = buscar_columna(df_status, ["Nro Solicitud", "Nro de Solicitud", "Nro. Solicitud"])
                                    status_nombre = buscar_columna(df_status, ["Nombre"])
                                    if not col_estado:
                                        st.error("El archivo de estado no contiene la columna 'Estado de OT'.")
                                        df_status = None
                                    else:
                                        with st.expander("👁️ Vista previa del estado de Laserfiche", expanded=False):
                                            st.dataframe(df_status.head(10), use_container_width=True, hide_index=True)
                                else:
                                    df_status, status_id = None, "SIN_ARCHIVO_ESTADO"

                                comparison_key = source_id + ":" + compare_id + ":" + status_id
                                if st.button("⚖️ Comparar y preparar resultado", type="primary", key="btn_comparar_items"):
                                    # Mapa código -> filas de documentos prestados; conserva todos los datos pedidos.
                                    prest = df_compare.copy()
                                    prest["__CODIGO_NORMALIZADO__"] = prest[col_codigo].map(normalizar_codigo)
                                    prest = prest[prest["__CODIGO_NORMALIZADO__"] != ""]
                                    base = resultado.copy()
                                    base["__CODIGO_NORMALIZADO__"] = base[items_resultado].map(normalizar_codigo)
                                    comp = base.merge(prest, on="__CODIGO_NORMALIZADO__", how="left", suffixes=("", " - Documento prestado"), indicator=True)
                                    comp["Estado comparación"] = comp["_merge"].map({"both": "DUPLICADO", "left_only": "NO ENCONTRADO"})
                                    comp.drop(columns=["_merge", "__CODIGO_NORMALIZADO__"], inplace=True, errors="ignore")
                                    # Asegura nombres de columnas claros y agrega los tres campos requeridos si el merge generó sufijos.
                                    for original, canonico in [(col_codigo, "Codigo del Item"), (col_solicitud, "Nro de Solicitud"), (col_fecha_solicitud, "Fecha de Solicitud")]:
                                        candidato = original
                                        if candidato not in comp.columns:
                                            candidato = original + " - Documento prestado"
                                        if candidato in comp.columns and canonico not in comp.columns:
                                            comp.rename(columns={candidato: canonico}, inplace=True)
                                    # Añade Estado de OT buscando primero por Nro Solicitud y luego por Nombre.
                                    comp["Estado de OT"] = ""
                                    if df_status is not None and col_estado:
                                        estado_map = {}
                                        if status_solicitud:
                                            for _, rr in df_status.iterrows():
                                                k = normalizar_codigo(rr[status_solicitud])
                                                if k and str(rr[col_estado]).strip(): estado_map[("SOL", k)] = str(rr[col_estado]).strip()
                                        nombre_laser = buscar_columna(comp, ["Nombre"])
                                        if status_nombre:
                                            for _, rr in df_status.iterrows():
                                                k = normalizar_codigo(rr[status_nombre])
                                                if k and str(rr[col_estado]).strip(): estado_map[("NOM", k)] = str(rr[col_estado]).strip()
                                        for ix, rr in comp.iterrows():
                                            sol = normalizar_codigo(rr.get("Nro de Solicitud", ""))
                                            nom = normalizar_codigo(rr.get(nombre_laser, "")) if nombre_laser else ""
                                            comp.at[ix, "Estado de OT"] = estado_map.get(("SOL", sol), estado_map.get(("NOM", nom), ""))
                                    # Orden de columnas prácticas: datos LASER, datos del préstamo, comparación y estado OT.
                                    preferidas = [c for c in [nombre_resultado, fecha_resultado, items_resultado, "Codigo del Item", "Nro de Solicitud", "Fecha de Solicitud", "Estado de OT", "Estado comparación"] if c in comp.columns]
                                    resto = [c for c in comp.columns if c not in preferidas]
                                    comp = comp[preferidas + resto]
                                    st.session_state["resultado_comparacion"] = comp
                                    st.session_state["comparacion_source_id"] = comparison_key

                                if (st.session_state.get("comparacion_source_id") == comparison_key and "resultado_comparacion" in st.session_state):
                                    comp = st.session_state["resultado_comparacion"].copy()
                                    duplicados = int((comp["Estado comparación"] == "DUPLICADO").sum())
                                    no_encontrados = int((comp["Estado comparación"] == "NO ENCONTRADO").sum())
                                    m1, m2, m3 = st.columns(3)
                                    m1.metric("Items separados", f"{len(comp):,}")
                                    m2.metric("Duplicados / encontrados", f"{duplicados:,}")
                                    m3.metric("No encontrados", f"{no_encontrados:,}")
                                    filtro = st.radio("Mostrar resultados", ["Todos", "Solo duplicados", "Solo no encontrados"], horizontal=True, key="filtro_resultado_comparacion")
                                    mostrar = comp
                                    if filtro == "Solo duplicados": mostrar = comp[comp["Estado comparación"] == "DUPLICADO"]
                                    elif filtro == "Solo no encontrados": mostrar = comp[comp["Estado comparación"] == "NO ENCONTRADO"]
                                    def resaltar_fila(fila):
                                        if fila.get("Estado comparación") == "DUPLICADO":
                                            return ["background-color: #fff2cc; color: #7f0000; font-weight: bold"] * len(fila)
                                        return [""] * len(fila)
                                    st.dataframe(mostrar.style.apply(resaltar_fila, axis=1), use_container_width=True, hide_index=True)
                                    d1, d2 = st.columns(2)
                                    with d1:
                                        st.download_button("📥 Descargar comparación completa", data=crear_excel(comp, "Comparacion", "Estado comparación"),
                                            file_name="COMPARACION_ITEMS_ONEIL_COMPLETA.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="descargar_comparacion_items")
                                    with d2:
                                        solo_dup = comp[comp["Estado comparación"] == "DUPLICADO"].copy()
                                        st.download_button("📥 Descargar solo duplicados", data=crear_excel(solo_dup, "Duplicados", "Estado comparación"),
                                            file_name="DUPLICADOS_ITEMS_ONEIL.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="descargar_solo_duplicados")
                                    if status_file is None:
                                        st.info("No se cargó el archivo de estado; la columna 'Estado de OT' queda vacía. Puedes subir stado-laser.xlsx y volver a comparar.")
                        except Exception as error:
                            st.error(f"No se pudo leer o comparar el archivo: {error}")

                    st.caption("Los archivos originales no se modifican. Los códigos se comparan como texto, ignorando espacios y diferencias entre mayúsculas/minúsculas.")
    except Exception as error:
        st.error(f"No se pudo leer o procesar el archivo: {error}")
